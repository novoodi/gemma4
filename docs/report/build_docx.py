"""
docs/report/part*.md → .docx 변환기 (python-docx).

한글(HWP)이 .docx를 직접 열 수 있으므로, 생성된 파일을 한글에서 불러와
중간 보고서 양식(표지·머리글)에 붙여 넣는 용도. 지원하는 마크다운 부분집합:
  #/##/###/#### 제목, 단락, '- ' 불릿, '1. ' 번호, '|' 표, ``` 코드블록, > 인용, **굵게**

사용법:
  python docs/report/build_docx.py                # 모든 part*.md → docs/report/out/*.docx
  python docs/report/build_docx.py part1_proposal # 특정 파일만
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = Path(__file__).parent
OUT = HERE / "out"

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")


def add_runs(paragraph, text):
    """**굵게** 와 `코드` 인라인 마크업을 run으로 분해."""
    pos = 0
    tokens = []
    for m in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
        if m.start() > pos:
            tokens.append(("plain", text[pos:m.start()]))
        if m.group(1) is not None:
            tokens.append(("bold", m.group(1)))
        else:
            tokens.append(("code", m.group(2)))
        pos = m.end()
    if pos < len(text):
        tokens.append(("plain", text[pos:]))
    for kind, t in tokens:
        run = paragraph.add_run(t)
        if kind == "bold":
            run.bold = True
        elif kind == "code":
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)


def set_cell_shading(cell, hex_fill):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)



LABEL_FILL = "E7E6E6"
FONT_SMALL = Pt(9)


def _clear_cell(cell):
    """cell.text = "" 는 빈 런을 남긴다(한글 변환기가 오해할 수 있음) — 런만 제거."""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    return p


def _para_fmt(p):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15


def _style_cell(cell, text, label=False, fs=FONT_SMALL):
    lines = text.replace("<br>", "\n").split("\n")
    p = _clear_cell(cell)
    for idx, line in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
        add_runs(p, line)
        _para_fmt(p)
        for r in p.runs:
            r.font.size = fs
            if label:
                r.bold = True
    if label:
        set_cell_shading(cell, LABEL_FILL)


def _set_widths(table, widths_in):
    """tblGrid 의 gridCol 과 각 셀의 tcW 를 같은 값으로 맞춘다(병합 셀은 span 합계).
    격자와 셀 너비가 어긋나면 한글(HWP) docx 변환기가 셀 영역을 잘못 그린다."""
    table.autofit = False
    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for gc, w in zip(grid_cols, widths_in):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        col = 0
        for tc in row._tr.tc_lst:
            span = tc.grid_span
            width = sum(widths_in[col:col + span])
            tc.width = Inches(width)
            col += span


def _grid(doc, ncols):
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    return table


def _row(table, cells):
    """cells: [(text, label, span), ...] — span 은 병합할 열 수."""
    row = table.add_row()
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    row.height = Cm(0.7)
    col = 0
    for text, label, span in cells:
        c = row.cells[col]
        if span > 1:
            c = c.merge(row.cells[col + span - 1])
        _style_cell(c, text, label)
        col += span
    return row


def _split2(value):
    parts = value.split(" / ", 1)
    return (parts[0].strip(), parts[1].strip() if len(parts) > 1 else "")


def _split3(value):
    parts = [v.strip() for v in value.split(" / ", 2)]
    while len(parts) < 3:
        parts.append("")
    return parts


def render_spec_grid(doc, rows, heading_name):
    """요구사항 항목표(UC/IFR/UIR/PER/SER) → 원본 4열 격자."""
    kv = {r[0].strip(): (r[1] if len(r) > 1 else "") for r in rows}
    num, typ = _split2(kv.get("요구사항 번호 / 유형", ""))
    name = kv.get("요구사항 이름", heading_name)
    W = [1.15, 2.0, 1.15, 2.0]
    t = _grid(doc, 4)
    if num.startswith("PER-"):
        _row(t, [("[" + num + "] " + name, True, 4)])
    _row(t, [("요구사항 번호", True, 1), (num, False, 1), ("요구사항 유형", True, 1), (typ, False, 1)])
    _row(t, [("요구사항 이름", True, 1), (name, False, 3)])
    _row(t, [("요구사항 개요", True, 1), (kv.get("요구사항 개요", ""), False, 3)])
    _row(t, [("요구사항 내용", True, 1), (kv.get("요구사항 내용", ""), False, 3)])
    if "입력 데이터 / 출력 데이터" in kv:
        i, o = _split2(kv["입력 데이터 / 출력 데이터"])
        _row(t, [("입력 데이터", True, 1), (i, False, 1), ("출력 데이터", True, 1), (o, False, 1)])
    if "입출력 유형 / 데이터 파일 유형" in kv:
        a, b = _split2(kv["입출력 유형 / 데이터 파일 유형"])
        _row(t, [("입출력 유형", True, 1), (a, False, 1), ("데이터 파일 유형", True, 1), (b, False, 1)])
    imp, risk = _split2(kv.get("중요도 / 위험도", ""))
    _row(t, [("중요도", True, 1), (imp, False, 1), ("위험도", True, 1), (risk, False, 1)])
    _row(t, [("품질속성", True, 1), (kv.get("품질속성", ""), False, 3)])
    _row(t, [("평가 방법", True, 1), (kv.get("평가 방법", ""), False, 3)])
    _row(t, [("관련 요구사항", True, 1), (kv.get("관련 요구사항", ""), False, 3)])
    _row(t, [("요구사항 출처", True, 1), (kv.get("요구사항 출처", ""), False, 3)])
    _set_widths(t, W)
    doc.add_paragraph()


def render_qr_grid(doc, header, rows):
    """품질 요구사항표 → 원본 4열 격자."""
    kv = {r[0].strip(): (r[1] if len(r) > 1 else "") for r in rows}
    title = header[1].strip()
    num, name = title.split(" ", 1) if " " in title else (title, "")
    W = [1.15, 2.0, 1.15, 2.0]
    t = _grid(doc, 4)
    _row(t, [("요구사항 번호", True, 1), (num, False, 1), ("요구사항 유형", True, 1), ("품질 (비기능)", False, 1)])
    _row(t, [("요구사항 이름", True, 1), (name, False, 3)])
    _row(t, [("품질 특성", True, 1), (kv.get("품질 특성", ""), False, 3)])
    if "요구사항 개요" in kv:
        _row(t, [("요구사항 개요", True, 1), (kv["요구사항 개요"], False, 3)])
    _row(t, [("요구사항 내용", True, 1), (kv.get("요구사항 내용", ""), False, 3)])
    _row(t, [("측정 지표 및 목표값", True, 1), (kv.get("측정 지표 및 목표값", ""), False, 3)])
    imp, risk = _split2(kv.get("중요도 / 위험도", ""))
    _row(t, [("중요도", True, 1), (imp, False, 1), ("위험도", True, 1), (risk, False, 1)])
    _row(t, [("평가 방법", True, 1), (kv.get("평가 방법", ""), False, 3)])
    _row(t, [("관련 요구사항", True, 1), (kv.get("관련 요구사항", ""), False, 3)])
    if "요구사항 출처" in kv:
        _row(t, [("요구사항 출처", True, 1), (kv["요구사항 출처"], False, 3)])
    _set_widths(t, W)
    doc.add_paragraph()


def render_uc_grid(doc, rows):
    """유스케이스 명세·시퀀스 머리표 → 원본 6열 격자."""
    kv_rows = [(r[0].strip(), (r[1] if len(r) > 1 else "")) for r in rows]
    kv = dict(kv_rows)
    W = [0.95, 1.15, 0.95, 1.15, 0.95, 1.15]
    t = _grid(doc, 6)
    _row(t, [("시스템명", True, 1), (kv.get("시스템명", ""), False, 3),
             ("서브시스템명", True, 1), (kv.get("서브시스템명", ""), False, 1)])
    stage, date, ver = _split3(kv.get("단계명 / 작성일자 / 버전", ""))
    _row(t, [("단계명", True, 1), (stage, False, 1), ("작성일자", True, 1), (date, False, 1),
             ("버전", True, 1), (ver, False, 1)])
    id_key = "UC-ID / 이름" if "UC-ID / 이름" in kv else "Sequence-ID / 이름"
    uid, uname = _split2(kv.get(id_key, ""))
    _row(t, [(id_key.split(" / ")[0], True, 1), (uid, False, 2), ("이 름", True, 1), (uname, False, 2)])
    author, adate = _split2(kv.get("작성자 / 작성일", ""))
    _row(t, [("작성자", True, 1), (author, False, 2), ("작성일", True, 1), (adate, False, 2)])
    handled = {"시스템명", "서브시스템명", "단계명 / 작성일자 / 버전", id_key, "작성자 / 작성일"}
    for k, v in kv_rows:
        if k in handled:
            continue
        _row(t, [(k, True, 1), (v, False, 5)])
    _set_widths(t, W)
    doc.add_paragraph()

def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        _clear_cell(cell)
        add_runs(cell.paragraphs[0], h)
        for r in cell.paragraphs[0].runs:
            r.bold = True
        set_cell_shading(cell, "E7E6E6")
    for row in body:
        cells = table.add_row().cells
        for i in range(len(header)):
            _style_cell(cells[i], row[i] if i < len(row) else "", label=False)
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Cm(0.6)
    doc.add_paragraph()


def split_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path: Path, docx_path: Path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10.5)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    para_buf = []
    last_heading_name = ""

    def flush_para():
        nonlocal para_buf
        if para_buf:
            p = doc.add_paragraph()
            add_runs(p, " ".join(s.strip() for s in para_buf))
            para_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        m = re.match(r"^\(그림 삽입:\s*`([^`]+)`\s*(?:—\s*(.*?))?\)$", stripped)
        if m:
            flush_para()
            img = HERE / "figures" / m.group(1)
            if img.exists():
                doc.add_picture(str(img), width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if m.group(2):
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = cap.add_run(f"[그림] {m.group(2)}")
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            else:
                p = doc.add_paragraph()
                add_runs(p, f"(그림 파일 없음: {m.group(1)})")
            i += 1
            continue

        if stripped.startswith("```"):
            flush_para()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9)
            i += 1
            continue

        if stripped.startswith("|"):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_table_row(lines[i])
                if not all(re.fullmatch(r":?-+:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                header = [c.strip() for c in rows[0]]
                first = rows[1][0].strip() if len(rows) > 1 else ""
                if header[:2] == ["항목", "내용"] and first.startswith("요구사항 번호"):
                    render_spec_grid(doc, rows[1:], last_heading_name)
                elif len(header) > 1 and header[0] == "항목" and header[1].startswith("QR-"):
                    render_qr_grid(doc, header, rows[1:])
                elif header[:2] == ["항목", "내용"] and first == "시스템명":
                    render_uc_grid(doc, rows[1:])
                else:
                    add_table(doc, rows)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            flush_para()
            level = len(m.group(1))
            hm = re.match(r"^(?:UC-\d{3}-\d{3}|IFR-\d{3}|UIR-\d{3}|PER-\d{3}|SER-\d{3})\s+(.*?)\s*(?:【.*?】)?\s*$", m.group(2))
            if hm:
                last_heading_name = hm.group(1)
            h = doc.add_paragraph()
            add_runs(h, m.group(2))
            size = {1: 16, 2: 14, 3: 12, 4: 11}.get(level, 11)
            for r in h.runs:
                r.bold = True
                r.font.size = Pt(size)
                r.font.color.rgb = RGBColor(0, 0, 0)
                r.font.name = "맑은 고딕"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.line_spacing = 1.15
            i += 1
            continue

        if stripped.startswith(">"):
            flush_para()
            p = doc.add_paragraph()
            add_runs(p, stripped.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            i += 1
            continue

        m = re.match(r"^(\s*)-\s+(.*)$", line)
        if m:
            flush_para()
            indent = len(m.group(1))
            p = doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
            add_runs(p, m.group(2))
            # 다음 줄이 들여쓰기된 연속 텍스트면 같은 항목에 이어 붙임
            while i + 1 < len(lines) and lines[i + 1].startswith("  ") and not re.match(r"^\s*(-|\d+\.)\s", lines[i + 1]) and lines[i + 1].strip():
                i += 1
                p.add_run("\n")
                add_runs(p, lines[i].strip())
            i += 1
            continue

        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            flush_para()
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            i += 1
            continue

        if not stripped:
            flush_para()
            i += 1
            continue

        para_buf.append(line)
        i += 1

    flush_para()
    for p in doc.paragraphs:
        if p.paragraph_format.line_spacing is None:
            p.paragraph_format.line_spacing = 1.15
    OUT.mkdir(exist_ok=True)
    doc.save(docx_path)


def main():
    targets = sys.argv[1:]
    files = sorted(HERE.glob("part*.md"))
    if targets:
        files = [f for f in files if f.stem in targets]
    if not files:
        print("변환할 part*.md 가 없습니다.")
        return
    for f in files:
        out = OUT / f"{f.stem}.docx"
        convert(f, out)
        print(f"{f.name} → {out.relative_to(HERE.parent.parent)}")


if __name__ == "__main__":
    main()
